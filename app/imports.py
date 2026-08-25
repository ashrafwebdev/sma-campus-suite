"""Bulk import of roster data from a single multi-sheet Excel workbook.

`build_template()` produces the .xlsx a school fills in; `import_workbook()`
reads one back and creates records through the same crud functions the UI
uses (so admission numbers, employee numbers etc. are generated the same
way). Each row is processed independently -- one bad row is recorded as an
error and skipped, the rest of the import still goes through.
"""

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.crud import academic, hr, student
from app.models.academic import SchoolClass, Section
from app.schemas.academic import SchoolClassCreate, SectionCreate, SubjectCreate
from app.schemas.hr import EmployeeCreate
from app.schemas.student import StudentCreate

SUBJECT_TYPE_BY_LABEL = {"core": 1, "elective": 2, "selective": 3}
RESIDENCY_BY_LABEL = {"day scholar": 1, "hosteller": 2, "hostel": 2}

_HEADER_FILL = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
_HEADER_FONT = Font(bold=True)

# (sheet name, [(header, column width), ...])
_SHEETS: list[tuple[str, list[tuple[str, int]]]] = [
    ("Classes", [("Name*", 20), ("Order", 10)]),
    ("Sections", [("Name*", 14), ("Class Name*", 20), ("Capacity", 10)]),
    ("Subjects", [("Name*", 20), ("Code*", 12), ("Class Name*", 20), ("Type (Core/Elective/Selective)", 26)]),
    (
        "Students",
        [
            ("Name*", 22),
            ("Class Name", 18),
            ("Section Name", 14),
            ("Phone", 15),
            ("Email", 26),
            ("Guardian Name", 22),
            ("Guardian Phone", 15),
            ("Address", 30),
            ("Residency (Day Scholar/Hosteller)", 22),
        ],
    ),
    ("Employees", [("Name*", 22), ("Designation*", 20), ("Phone", 15), ("Email", 26), ("Basic Salary", 14)]),
]

_READ_ME = """Campus Suite -- Bulk Import Template

One Excel file, one sheet per record type. Fill in the sheets you need and
leave the rest empty -- empty sheets are skipped.

Columns marked with * are required. Everything else is optional.

Import order matters for lookups: Classes, then Sections, then Subjects,
then Students, then Employees. If you're adding students to classes that
already exist in the system, you can leave the Classes/Sections/Subjects
sheets empty and just fill in Students -- Class Name and Section Name are
matched against classes/sections already in the system too, not only ones
in this file.

Sheet: Classes
  Name* -- e.g. "Class 9"
  Order -- controls display order (lower first); defaults to 0

Sheet: Sections
  Name* -- e.g. "A"
  Class Name* -- must match a class name (from the Classes sheet above, or
    already in the system)
  Capacity -- defaults to 40

Sheet: Subjects
  Name* -- e.g. "Mathematics"
  Code* -- e.g. "MATH9"
  Class Name* -- must match a class name
  Type -- Core, Elective, or Selective; defaults to Core

Sheet: Students
  Name* -- required
  Class Name / Section Name -- optional; must match existing names if given
  Residency -- "Day Scholar" or "Hosteller"; defaults to Day Scholar
  Everything else is optional contact/guardian info

Sheet: Employees
  Name*, Designation* -- required
  Basic Salary -- plain number, e.g. 30000

If a row fails (e.g. an unknown class name, a missing required field), it's
skipped and reported after import -- the rest of the file still goes
through.
"""


@dataclass
class ImportResult:
    counts: dict[str, int] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)


def build_template() -> bytes:
    wb = Workbook()
    readme_ws = wb.active
    readme_ws.title = "Read Me"
    readme_ws.column_dimensions["A"].width = 90
    for i, line in enumerate(_READ_ME.strip("\n").split("\n"), start=1):
        cell = readme_ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=False)
        if line.startswith("Campus Suite") or line.startswith("Sheet:"):
            cell.font = Font(bold=True)

    for sheet_name, columns in _SHEETS:
        ws: Worksheet = wb.create_sheet(sheet_name)
        for col_idx, (header, width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            ws.column_dimensions[cell.column_letter].width = width
        ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _norm(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_rows(ws: Worksheet) -> list[dict[str, str]]:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    headers = [_norm(h).rstrip("*").strip().lower() for h in header_row]
    rows = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None or _norm(v) == "" for v in raw_row):
            continue
        row = {headers[i]: raw_row[i] for i in range(min(len(headers), len(raw_row)))}
        rows.append(row)
    return rows


def _to_int(value, default: int) -> int:
    text = _norm(value)
    if not text:
        return default
    return int(float(text))


def _to_decimal(value, default: str) -> Decimal:
    text = _norm(value)
    if not text:
        return Decimal(default)
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"'{text}' is not a valid number") from exc


def import_workbook(db: Session, file_bytes: bytes) -> ImportResult:
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read the file as an Excel workbook: {exc}") from exc

    result = ImportResult()

    def bump(table: str) -> None:
        result.counts[table] = result.counts.get(table, 0) + 1

    def get_class(name: str) -> SchoolClass | None:
        name = name.strip()
        return db.query(SchoolClass).filter(func.lower(SchoolClass.name) == name.lower()).first()

    def get_section(class_id: int, name: str) -> Section | None:
        name = name.strip()
        return (
            db.query(Section)
            .filter(Section.class_id == class_id, func.lower(Section.name) == name.lower())
            .first()
        )

    # -- Classes -------------------------------------------------------
    if "Classes" in wb.sheetnames:
        for i, row in enumerate(_sheet_rows(wb["Classes"]), start=2):
            name = _norm(row.get("name"))
            if not name:
                result.errors.append(f"Classes row {i}: Name is required")
                continue
            try:
                if get_class(name):
                    result.errors.append(f"Classes row {i}: a class named '{name}' already exists, skipped")
                    continue
                academic.create_class(db, SchoolClassCreate(name=name, order=_to_int(row.get("order"), 0)))
                bump("school_classes")
            except Exception as exc:
                db.rollback()
                result.errors.append(f"Classes row {i}: {exc}")

    # -- Sections --------------------------------------------------------
    if "Sections" in wb.sheetnames:
        for i, row in enumerate(_sheet_rows(wb["Sections"]), start=2):
            name = _norm(row.get("name"))
            class_name = _norm(row.get("class name"))
            if not name or not class_name:
                result.errors.append(f"Sections row {i}: Name and Class Name are both required")
                continue
            try:
                school_class = get_class(class_name)
                if not school_class:
                    result.errors.append(f"Sections row {i}: no class named '{class_name}' found")
                    continue
                if get_section(school_class.id, name):
                    result.errors.append(f"Sections row {i}: section '{name}' already exists in '{class_name}', skipped")
                    continue
                academic.create_section(
                    db, SectionCreate(name=name, capacity=_to_int(row.get("capacity"), 40), class_id=school_class.id)
                )
                bump("sections")
            except Exception as exc:
                db.rollback()
                result.errors.append(f"Sections row {i}: {exc}")

    # -- Subjects --------------------------------------------------------
    if "Subjects" in wb.sheetnames:
        for i, row in enumerate(_sheet_rows(wb["Subjects"]), start=2):
            name = _norm(row.get("name"))
            code = _norm(row.get("code"))
            class_name = _norm(row.get("class name"))
            if not name or not code or not class_name:
                result.errors.append(f"Subjects row {i}: Name, Code and Class Name are all required")
                continue
            try:
                school_class = get_class(class_name)
                if not school_class:
                    result.errors.append(f"Subjects row {i}: no class named '{class_name}' found")
                    continue
                type_label = _norm(row.get("type (core/elective/selective)") or row.get("type")).lower()
                subject_type = SUBJECT_TYPE_BY_LABEL.get(type_label, 1)
                academic.create_subject(
                    db, SubjectCreate(name=name, code=code, subject_type=subject_type, class_id=school_class.id)
                )
                bump("subjects")
            except Exception as exc:
                db.rollback()
                result.errors.append(f"Subjects row {i}: {exc}")

    # -- Students ----------------------------------------------------------
    if "Students" in wb.sheetnames:
        for i, row in enumerate(_sheet_rows(wb["Students"]), start=2):
            name = _norm(row.get("name"))
            if not name:
                result.errors.append(f"Students row {i}: Name is required")
                continue
            try:
                class_id = None
                section_id = None
                class_name = _norm(row.get("class name"))
                if class_name:
                    school_class = get_class(class_name)
                    if not school_class:
                        result.errors.append(f"Students row {i}: no class named '{class_name}' found")
                        continue
                    class_id = school_class.id
                    section_name = _norm(row.get("section name"))
                    if section_name:
                        section = get_section(class_id, section_name)
                        if not section:
                            result.errors.append(
                                f"Students row {i}: no section named '{section_name}' found in '{class_name}'"
                            )
                            continue
                        section_id = section.id

                residency_label = _norm(row.get("residency (day scholar/hosteller)") or row.get("residency")).lower()
                residency_type = RESIDENCY_BY_LABEL.get(residency_label, 1)

                student.create_student(
                    db,
                    StudentCreate(
                        name=name,
                        phone_no=_norm(row.get("phone")) or None,
                        email=_norm(row.get("email")) or None,
                        guardian_name=_norm(row.get("guardian name")) or None,
                        guardian_phone_no=_norm(row.get("guardian phone")) or None,
                        permanent_address=_norm(row.get("address")) or None,
                        residency_type=residency_type,
                        class_id=class_id,
                        section_id=section_id,
                    ),
                )
                bump("students")
            except Exception as exc:
                db.rollback()
                result.errors.append(f"Students row {i}: {exc}")

    # -- Employees -----------------------------------------------------
    if "Employees" in wb.sheetnames:
        for i, row in enumerate(_sheet_rows(wb["Employees"]), start=2):
            name = _norm(row.get("name"))
            designation = _norm(row.get("designation"))
            if not name or not designation:
                result.errors.append(f"Employees row {i}: Name and Designation are both required")
                continue
            try:
                hr.create_employee(
                    db,
                    EmployeeCreate(
                        name=name,
                        designation=designation,
                        phone_no=_norm(row.get("phone")) or None,
                        email=_norm(row.get("email")) or None,
                        basic_salary=_to_decimal(row.get("basic salary"), "0"),
                    ),
                )
                bump("employees")
            except Exception as exc:
                db.rollback()
                result.errors.append(f"Employees row {i}: {exc}")

    return result
