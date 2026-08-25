from io import BytesIO

from openpyxl import Workbook, load_workbook


def _workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(client, auth_headers, content: bytes, filename: str = "import.xlsx"):
    return client.post(
        "/api/v1/import",
        headers=auth_headers,
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_template_download_has_expected_sheets(client, auth_headers):
    response = client.get("/api/v1/import/template", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Read Me", "Classes", "Sections", "Subjects", "Students", "Employees"]
    assert wb["Classes"]["A1"].value == "Name*"
    assert wb["Students"]["A1"].value == "Name*"


def test_template_download_requires_auth(client):
    assert client.get("/api/v1/import/template").status_code == 401


def test_import_creates_records_across_all_sheets(client, auth_headers):
    content = _workbook_bytes(
        {
            "Classes": [["Name*", "Order"], ["Class 9", 1], ["Class 10", 2]],
            "Sections": [["Name*", "Class Name*", "Capacity"], ["A", "Class 9", 30]],
            "Subjects": [
                ["Name*", "Code*", "Class Name*", "Type"],
                ["Mathematics", "MATH9", "Class 9", "Core"],
            ],
            "Students": [
                ["Name*", "Class Name", "Section Name", "Phone", "Residency"],
                ["Asha Rao", "Class 9", "A", "9998887777", "Hosteller"],
                ["Bilal Khan", "Class 10", "", "", "Day Scholar"],
            ],
            "Employees": [
                ["Name*", "Designation*", "Basic Salary"],
                ["Nina Fernandes", "Teacher", 32000],
            ],
        }
    )

    response = _upload(client, auth_headers, content)
    assert response.status_code == 200
    body = response.json()
    assert body["errors"] == []
    assert body["counts"] == {
        "school_classes": 2,
        "sections": 1,
        "subjects": 1,
        "students": 2,
        "employees": 1,
    }

    classes = client.get("/api/v1/academic/classes", headers=auth_headers).json()
    class_names = {c["name"] for c in classes}
    assert {"Class 9", "Class 10"} <= class_names

    students = client.get("/api/v1/students", headers=auth_headers).json()
    asha = next(s for s in students if s["name"] == "Asha Rao")
    assert asha["residency_type"] == 2
    class9_id = next(c["id"] for c in classes if c["name"] == "Class 9")
    assert asha["class_id"] == class9_id

    bilal = next(s for s in students if s["name"] == "Bilal Khan")
    assert bilal["residency_type"] == 1
    assert bilal["section_id"] is None

    employees = client.get("/api/v1/hr/employees", headers=auth_headers).json()
    assert any(e["name"] == "Nina Fernandes" and e["basic_salary"] == "32000.00" for e in employees)


def test_import_resolves_class_already_in_the_system(client, auth_headers):
    existing_class = client.post("/api/v1/academic/classes", json={"name": "Grade 5"}, headers=auth_headers).json()

    content = _workbook_bytes(
        {"Students": [["Name*", "Class Name"], ["Priya Nair", "Grade 5"]]},
    )
    response = _upload(client, auth_headers, content)
    assert response.status_code == 200
    assert response.json()["errors"] == []

    students = client.get("/api/v1/students", headers=auth_headers).json()
    priya = next(s for s in students if s["name"] == "Priya Nair")
    assert priya["class_id"] == existing_class["id"]


def test_import_reports_row_errors_without_failing_whole_file(client, auth_headers):
    content = _workbook_bytes(
        {
            "Classes": [["Name*", "Order"], ["", 1], ["Class 11", 1]],
            "Students": [["Name*", "Class Name"], ["Orphan Student", "No Such Class"]],
        }
    )
    response = _upload(client, auth_headers, content)
    assert response.status_code == 200
    body = response.json()
    assert body["counts"] == {"school_classes": 1}
    assert len(body["errors"]) == 2
    assert any("Classes row 2" in e for e in body["errors"])
    assert any("Students row 2" in e and "No Such Class" in e for e in body["errors"])

    classes = client.get("/api/v1/academic/classes", headers=auth_headers).json()
    assert any(c["name"] == "Class 11" for c in classes)
    assert client.get("/api/v1/students", headers=auth_headers).json() == []


def test_import_rejects_non_xlsx_file(client, auth_headers):
    response = client.post(
        "/api/v1/import",
        headers=auth_headers,
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert response.status_code == 400


def test_import_requires_auth(client):
    response = client.post("/api/v1/import", files={"file": ("import.xlsx", b"", "application/octet-stream")})
    assert response.status_code == 401
